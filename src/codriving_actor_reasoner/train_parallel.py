import os
os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"
import ssl
orig_load_default_certs = ssl.SSLContext.load_default_certs
def patched_load_default_certs(self, purpose=ssl.Purpose.SERVER_AUTH):
    try:
        orig_load_default_certs(self, purpose)
    except Exception:
        pass
ssl.SSLContext.load_default_certs = patched_load_default_certs

import time
import argparse
import imageio
import openpyxl
import gym
import numpy as np
import re
from openai import OpenAI
from llm_controller.memory import DrivingMemory
from llm_controller.parallel_agent import ParallelAgentCoordination
from experiment_scenarios import create_environment

def open_excel(i, scenario_name):
    file_dir = f'./results/train/{scenario_name}/excel/'
    file_name = file_dir + str(i) + '.xlsx'

    if not os.path.exists(file_dir):
        os.makedirs(file_dir)
    workbook = openpyxl.Workbook()
    if os.path.exists(file_name):
        workbook = openpyxl.load_workbook(file_name)
    return file_name, workbook

def clean_val(val):
    if isinstance(val, str):
        val = val.replace('\r', ' ').replace('\n', ' ')
        val = re.sub(r'[\000-\010]|[\013-\014]|[\016-\037]', '', val)
    return val

def write_data(workbook, env, t):
    column_names = ['t', 'x', 'y', 'v', 'theta', 'background_veh?']
    for vehicle in env.road.vehicles:
        sheet_name = str(vehicle.id)
        if sheet_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(sheet_name)
            worksheet.append(column_names)
        else:
            worksheet = workbook[sheet_name]
        controlled_vehicles = env.controlled_vehicles
        background_vehicles = vehicle not in controlled_vehicles
        state = [round(vehicle.position[0], 2), round(vehicle.position[1], 2), round(vehicle.speed, 2), round(vehicle.heading, 2), background_vehicles]
        row_data = [t, round(vehicle.position[0], 2), round(vehicle.position[1], 2), round(vehicle.speed, 2), round(vehicle.heading, 2), background_vehicles]
        
        state = [clean_val(item) for item in state]
        row_data = [clean_val(item) for item in row_data]
        
        worksheet.append(row_data)
        worksheet.cell(row=t + 2, column=1, value=t)
        for i, item in enumerate(state):
            worksheet.cell(row=t + 2, column=i + 2, value=item)
    return workbook

def validate_dispatched_action(vehicle, action_trace, dispatched_action):
    if dispatched_action != action_trace['final_action_id']:
        return False, 'post-shield action was not dispatched'
    if vehicle.crashed:
        return False, 'vehicle crashed before completing 2-step horizon'
    if not vehicle.on_road:
        return False, 'vehicle went off-road before completing 2-step horizon'
    return True, None

def log_validation_status(transition, status, reason=None):
    message = (
        f"[Validated Memory] CAV {transition['vehicle_id']} | "
        f"Origin step: {transition['origin_step']} | "
        f"Executed: {transition['final_action']} | "
        f"Status: {status}"
    )
    if reason:
        message += f" | Reason: {reason}"
    print(message)

def write_validated_transition(memory, action_module, transition):
    executed_action_val = [
        np.array([transition['final_action_id']]),
        transition['style'],
        transition['intention'],
    ]
    action_module.memory_update(
        memory,
        transition['prompt_info'],
        executed_action_val
    )

# Set up environment (enable merge or intersection or highway)
parser = argparse.ArgumentParser(description="Run parallel-agent training.")
parser.add_argument(
    "--scenario",
    choices=["highway", "intersection", "merge"],
    default="highway",
    help="Traffic scenario to run (default: highway)."
)
parser.add_argument(
    "--episodes",
    type=int,
    default=50,
    help="Number of training episodes to run (default: 50)."
)
args = parser.parse_args()

env = create_environment(gym, args.scenario)
scenario_name = env.spec.id

print(f"Requested scenario: {args.scenario}")
print(f"Actual environment: {env.spec.id}")
print(f"Controlled vehicles: {len(env.controlled_vehicles)}")
print(f"Environment action type: {type(env.action_type).__name__}")
print(f"Episodes: {args.episodes}")

# Enable training mode (with online memory writing)
os.environ["EMBEDDING_PROVIDER"] = "ollama"  # set to ollama or openai
os.environ["LLM_PROVIDER"] = "ollama"

client = OpenAI(
    api_key=os.getenv("LLM_API_KEY"),
    base_url=os.getenv("LLM_API_BASE", "http://127.0.0.1:11434/v1")
)
model_name = os.getenv("LLM_MODEL", "qwen2.5:7b")

episodes = args.episodes  # train mode defaults to 50 rounds for learning

print("==================== Starting Parallel Agent Training (Online Learning) ====================")

for i in range(episodes):
    video_path = f'./results/train/{scenario_name}/video/{i}.mp4'
    os.makedirs(os.path.dirname(video_path), exist_ok=True)
    writer = imageio.get_writer(video_path, fps=30)
    file_name, workbook = open_excel(i, scenario_name)
    terminated = False
    t = 0
    obs = env.reset()
    
    memory = DrivingMemory(env)
    coordinator = ParallelAgentCoordination(env, memory)
    coordinator.start_reasoner(env)
    
    # Store history of prompts and actions to allow self-reflection upon collision
    episode_history = []
    pending_transitions = []
    
    while not terminated:
        print('---------------------------------------------------------------')
        print(f"Episode {i} | Step {t} | Training Mode")
        
        # 1. Fast RAG retrieval and safety verification
        llm_actions = coordinator.get_actor_actions(env, memory)
        action = [item for sublist in llm_actions for item in sublist]
        action_trace = coordinator.get_last_actor_action_trace()
        
        # Record prompt context and action taken
        latest_negotiation_results = coordinator.action_module.transfer_negotiation_prompts_to_results(
            env.controlled_vehicles[0], coordinator.shared_negotiation_prompt
        )
        latest_prompt_info = coordinator.action_module.prompt_engineer(
            env.controlled_vehicles[0], env.road, env, latest_negotiation_results, coordinator.shared_conflicting_info
        )
        
        episode_history.append({
            'prompt_info': latest_prompt_info,
            'action_taken': action,
            'negotiation': latest_negotiation_results
        })
        
        obs, global_reward, terminated, info = env.step(tuple(action), env)

        # First give only transitions from earlier steps their next observation.
        # This snapshot prevents a new transition from being observed twice here.
        older_pending = pending_transitions
        pending_transitions = []
        for transition in older_pending:
            vehicle_index = transition['vehicle_index']
            if vehicle_index >= len(env.controlled_vehicles):
                log_validation_status(
                    transition,
                    'REJECTED',
                    'controlled-vehicle index no longer exists'
                )
                continue

            vehicle = env.controlled_vehicles[vehicle_index]
            same_vehicle = getattr(vehicle, 'id', id(vehicle)) == transition['vehicle_id']
            if not same_vehicle:
                log_validation_status(
                    transition,
                    'REJECTED',
                    'controlled-vehicle identity changed'
                )
                continue

            survived, failure_reason = validate_dispatched_action(
                vehicle, transition, transition['final_action_id']
            )
            if not survived:
                log_validation_status(transition, 'REJECTED', failure_reason)
                continue

            transition['survival_count'] += 1
            if transition['survival_count'] == 2:
                write_validated_transition(
                    memory, coordinator.action_module, transition
                )
                log_validation_status(transition, 'ACCEPTED (2/2)')
            else:
                pending_transitions.append(transition)
                log_validation_status(
                    transition,
                    f"PENDING ({transition['survival_count']}/2)"
                )

        # Then create transitions for actions executed by this env.step(). Each
        # new transition receives only its first post-step survival observation.
        dispatch_aligned = len(action) == len(action_trace) == len(env.controlled_vehicles)
        for vehicle_index, vehicle in enumerate(env.controlled_vehicles):
            if dispatch_aligned:
                transition = dict(action_trace[vehicle_index])
                transition['vehicle_index'] = vehicle_index
                transition['origin_step'] = t
                transition['survival_count'] = 0
                dispatched_action = action[vehicle_index]
                survived, validation_reason = validate_dispatched_action(
                    vehicle, transition, dispatched_action
                )
            else:
                transition = dict(action_trace[vehicle_index]) if vehicle_index < len(action_trace) else {
                    'vehicle_id': getattr(vehicle, 'id', id(vehicle)),
                    'actor_selected_action': 'UNKNOWN',
                    'final_action': 'UNKNOWN',
                    'final_action_id': None,
                    'shield_overrode': False,
                }
                transition['vehicle_index'] = vehicle_index
                transition['origin_step'] = t
                transition['survival_count'] = 0
                survived = False
                validation_reason = 'CAV/action trace alignment mismatch'

            if survived:
                transition['survival_count'] = 1
                pending_transitions.append(transition)
                log_validation_status(transition, 'PENDING (1/2)')
            else:
                log_validation_status(
                    transition, 'REJECTED', validation_reason
                )

        env.render()
        
        print("Action executed:", action)
        print("Reward:", global_reward)
        
        frame = env.render('rgb_array')
        writer.append_data(frame)
        
        workbook = write_data(workbook, env, t)
        workbook.save(file_name)
        t += 1
        
        # Check for collision
        is_collision = info.get("cav_crashed", False)
        if terminated or is_collision:
            for transition in pending_transitions:
                log_validation_status(
                    transition,
                    'CENSORED',
                    'episode ended before completing validation horizon'
                )
            pending_transitions.clear()

        if is_collision:
            print("❌ COLLISION DETECTED! Triggering Self-Reflection Refinement Loop...")
            
            # Retrieve last state before collision
            if len(episode_history) > 0:
                last_state = episode_history[-1]
                prompt_info = last_state['prompt_info']
                action_taken = last_state['action_taken']
                
                # Call LLM to reflect on why it crashed and suggest a correction
                reflection_prompt = (
                    "You are a driving expert analyzing a crash in a simulator.\n"
                    "Here is the scenario context right before the crash:\n"
                    f"{prompt_info}\n\n"
                    f"Ego vehicle took action index {action_taken} and CRASHED.\n"
                    "Analyze why this action led to a crash, and suggest the CORRECT action (choose from: FASTER, SLOWER, IDLE, LANE_LEFT, LANE_RIGHT).\n"
                    "Format your output exactly as a JSON block with 'correct_action' and 'reasoning' fields:\n"
                    "{\n"
                    "  \"correct_action\": \"<action_name>\",\n"
                    "  \"reasoning\": \"<brief explanation of why and what to avoid>\"\n"
                    "}"
                )
                
                try:
                    completion = client.chat.completions.create(
                        model=model_name,
                        messages=[{"role": "system", "content": reflection_prompt}],
                        temperature=0.0
                    )
                    response = completion.choices[0].message.content
                    print(f"Self-Reflection response:\n{response}")
                    
                    # Extract correct action and reasoning
                    match_act = re.search(r'"correct_action":\s*"(\w+)"', response)
                    match_reason = re.search(r'"reasoning":\s*"([^"]+)"', response)
                    
                    if match_act:
                        correct_action = match_act.group(1).upper()
                        comment = match_reason.group(1) if match_reason else "avoid collision"
                        print(
                            f"Self-Correction generated for analysis only: "
                            f"{correct_action} ({comment}). Not added to validated Actor memory."
                        )
                except Exception as ex:
                    print(f"Failed to complete self-reflection: {ex}")
            break
            
        time.sleep(0.1)
        
    coordinator.stop_reasoner()
    writer.close()
    print(f"Finished episode: {i}\n")
    
print("==================== Parallel Agent Training Complete ====================")
