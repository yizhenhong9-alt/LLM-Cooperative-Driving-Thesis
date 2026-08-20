import os
os.environ["KMP_DUPLICATE_LIB_OK"] = "TRUE"
import ssl
orig_load_default_certs = ssl.SSLContext.load_default_certs
def patched_load_default_certs(self, purpose=ssl.Purpose.SERVER_AUTH):
    try:
        orig_load_default_certs(self, purpose)
    except Exception:
        pass
ssl.SSLContext.load_default_certs = patched_load_default_certs

import time
import imageio
import openpyxl
import gym
import numpy as np
import re
from llm_controller.memory import DrivingMemory
from llm_controller.parallel_agent import ParallelAgentCoordination

def open_excel(i, scenario_name):
    file_dir = f'./results/test/{scenario_name}/excel/'
    file_name = file_dir + str(i) + '.xlsx'

    if not os.path.exists(file_dir):
        os.makedirs(file_dir)
    workbook = openpyxl.Workbook()
    if os.path.exists(file_name):
        workbook = openpyxl.load_workbook(file_name)
    return file_name, workbook

def clean_val(val):
    if isinstance(val, str):
        val = val.replace('\r', ' ').replace('\n', ' ')
        val = re.sub(r'[\000-\010]|[\013-\014]|[\016-\037]', '', val)
    return val

def write_data(workbook, env, t):
    column_names = ['t', 'x', 'y', 'v', 'theta', 'background_veh?']
    for vehicle in env.road.vehicles:
        sheet_name = str(vehicle.id)
        if sheet_name not in workbook.sheetnames:
            worksheet = workbook.create_sheet(sheet_name)
            worksheet.append(column_names)
        else:
            worksheet = workbook[sheet_name]
        controlled_vehicles = env.controlled_vehicles
        background_vehicles = vehicle not in controlled_vehicles
        state = [round(vehicle.position[0], 2), round(vehicle.position[1], 2), round(vehicle.speed, 2), round(vehicle.heading, 2), background_vehicles]
        row_data = [t, round(vehicle.position[0], 2), round(vehicle.position[1], 2), round(vehicle.speed, 2), round(vehicle.heading, 2), background_vehicles]
        
        state = [clean_val(item) for item in state]
        row_data = [clean_val(item) for item in row_data]
        
        worksheet.append(row_data)
        worksheet.cell(row=t + 2, column=1, value=t)
        for i, item in enumerate(state):
            worksheet.cell(row=t + 2, column=i + 2, value=item)
    return workbook

# Set up environment
env = gym.make('highway-v0')
scenario_name = env.spec.id

# Configure API details
os.environ["EMBEDDING_PROVIDER"] = "ollama"  # set to ollama or openai
os.environ["LLM_PROVIDER"] = "ollama"

# Test mode loops through 100 episodes
episodes = 100
success_count = 0
collisions = 0
speeds = []

print("==================== Starting Parallel Agent Testing (Evaluation Mode) ====================")

for i in range(episodes):
    video_path = f'./results/test/{scenario_name}/video/{i}.mp4'
    os.makedirs(os.path.dirname(video_path), exist_ok=True)
    writer = imageio.get_writer(video_path, fps=30)
    file_name, workbook = open_excel(i, scenario_name)
    terminated = False
    t = 0
    
    # Reset environment with deterministic seed (0 to 99)
    try:
        obs = env.reset(seed=i)
    except TypeError:
        import random
        random.seed(i)
        np.random.seed(i)
        if hasattr(env, 'seed') and callable(env.seed):
            env.seed(i)
        elif hasattr(env, 'unwrapped') and hasattr(env.unwrapped, 'seed'):
            if callable(env.unwrapped.seed):
                env.unwrapped.seed(i)
            else:
                env.unwrapped.seed = i
        obs = env.reset()
        
    memory = DrivingMemory(env)
    coordinator = ParallelAgentCoordination(env, memory)
    
    # Set training flag to False (Read-only database queries)
    coordinator.is_training = False
    
    # Start the background Reasoner thread
    coordinator.start_reasoner(env)
    
    episode_speeds = []
    crashed = False
    
    while not terminated:
        print('---------------------------------------------------------------')
        print(f"Episode {i} | Step {t} | Evaluation Mode")
        
        # 1. Fast retrieval via Actor memory cache & safety checks
        llm_actions = coordinator.get_actor_actions(env, memory)
        action = [item for sublist in llm_actions for item in sublist]
        
        # Capture current speed of the ego CAV
        ego = env.controlled_vehicles[0]
        episode_speeds.append(ego.speed)
        
        obs, global_reward, terminated, info = env.step(tuple(action), env)
        env.render()
        
        print("Action executed:", action)
        print("Ego speed:", round(ego.speed, 2))
        
        frame = env.render('rgb_array')
        writer.append_data(frame)
        
        workbook = write_data(workbook, env, t)
        workbook.save(file_name)
        t += 1
        
        # Check for collision
        if info.get("crashed", False):
            crashed = True
            print("💥 COLLISION occurred!")
            break
            
        time.sleep(0.1)
        
    # Stop background Reasoner
    coordinator.stop_reasoner()
    writer.close()
    
    if not crashed:
        success_count += 1
        avg_ep_speed = np.mean(episode_speeds) if episode_speeds else 0
        speeds.append(avg_ep_speed)
        print(f"✅ Episode {i} finished successfully. Avg speed: {round(avg_ep_speed, 2)} m/s")
    else:
        collisions += 1
        print(f"❌ Episode {i} failed due to collision.")
        
    print(f"Current stats after {i+1} rounds: Success rate = {round(success_count/(i+1)*100, 1)}%, Collisions = {collisions}\n")

# Output final statistical summary
print("==================== Evaluation Summary ====================")
print(f"Total episodes tested: {episodes}")
print(f"Success rate: {round(success_count/episodes*100, 1)}%")
print(f"Collision rate: {round(collisions/episodes*100, 1)}%")
if speeds:
    print(f"Average vehicle speed across success runs: {round(np.mean(speeds), 2)} m/s")
print("=============================================================")
