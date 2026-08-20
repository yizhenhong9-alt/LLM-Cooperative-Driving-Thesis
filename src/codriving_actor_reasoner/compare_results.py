import os
import glob
from openpyxl import load_workbook

COLLISION_THRESHOLD = 3.0  # distance threshold in meters for collision
INTERSECTION_CENTER_X = 0.0
INTERSECTION_CENTER_Y = 0.0
INTERSECTION_RADIUS = 8.0  # conflict zone radius in meters

def analyze_folder(folder_path):
    excel_path = os.path.join(folder_path, "*.xlsx")
    files = glob.glob(excel_path)
    
    if not files:
        return None
        
    successful_rounds = 0
    total_rounds = len(files)
    cav_speeds = []
    all_pets = []
    
    for file in files:
        try:
            # Open workbook in read-only and data-only mode for maximum speed
            wb = load_workbook(file, read_only=True, data_only=True)
            
            # Map time step t to active vehicles at that step
            active_vehs = {}
            # Record conflict intervals for PET calculation
            conflict_intervals = []
            
            max_t = 0
            
            for sheet_name in wb.sheetnames:
                if sheet_name.lower() == 'sheet':
                    continue
                
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows(values_only=True))
                if len(rows) <= 1:
                    continue
                
                header = rows[0]
                # Map headers to indices
                try:
                    t_idx = header.index('t')
                    x_idx = header.index('x')
                    y_idx = header.index('y')
                    v_idx = header.index('v')
                    bg_idx = header.index('background_veh?')
                except ValueError:
                    continue
                
                in_conflict = False
                entry_t = None
                
                for row in rows[1:]:
                    if row[t_idx] is None or row[x_idx] is None or row[y_idx] is None or row[v_idx] is None:
                        continue
                    
                    t_val = int(row[t_idx])
                    x = float(row[x_idx])
                    y = float(row[y_idx])
                    v = float(row[v_idx])
                    is_bg = bool(row[bg_idx])
                    
                    max_t = max(max_t, t_val)
                    
                    if t_val not in active_vehs:
                        active_vehs[t_val] = []
                    active_vehs[t_val].append((sheet_name, x, y, v, is_bg))
                    
                    # PET conflict calculation
                    dist_to_center = ((x - INTERSECTION_CENTER_X)**2 + (y - INTERSECTION_CENTER_Y)**2)**0.5
                    if dist_to_center < INTERSECTION_RADIUS:
                        if not in_conflict:
                            in_conflict = True
                            entry_t = t_val
                    else:
                        if in_conflict:
                            in_conflict = False
                            conflict_intervals.append((sheet_name, entry_t, t_val))
                            entry_t = None
                
                if in_conflict and entry_t is not None:
                    conflict_intervals.append((sheet_name, entry_t, max_t))
            
            # 1. Collision check
            is_collision = False
            for t_val, vehs in active_vehs.items():
                for idx1 in range(len(vehs)):
                    for idx2 in range(idx1 + 1, len(vehs)):
                        name1, x1, y1, v1, bg1 = vehs[idx1]
                        name2, x2, y2, v2, bg2 = vehs[idx2]
                        
                        dist = ((x1 - x2)**2 + (y1 - y2)**2)**0.5
                        if dist < COLLISION_THRESHOLD:
                            is_collision = True
                            break
                    if is_collision:
                        break
                if is_collision:
                    break
            
            if not is_collision:
                successful_rounds += 1
                
                # 2. Gather CAV speeds (where is_bg is False)
                for t_val, vehs in active_vehs.items():
                    for name, x, y, v, is_bg in vehs:
                        if not is_bg:
                            cav_speeds.append(v)
                            
                # 3. Process PET
                for idx1 in range(len(conflict_intervals)):
                    for idx2 in range(idx1 + 1, len(conflict_intervals)):
                        name1, entry1, exit1 = conflict_intervals[idx1]
                        name2, entry2, exit2 = conflict_intervals[idx2]
                        if name1 == name2:
                            continue
                        if entry1 < entry2:
                            pet = (entry2 - exit1) * 0.1
                        else:
                            pet = (entry1 - exit2) * 0.1
                        if pet > 0:
                            all_pets.append(pet)
                            
            wb.close()
        except Exception as e:
            # print(f"Error reading {file}: {e}")
            pass

    success_rate = (successful_rounds / total_rounds) * 100.0 if total_rounds > 0 else 0
    avg_cav_speed = sum(cav_speeds) / len(cav_speeds) if cav_speeds else 0
    avg_pet = sum(all_pets) / len(all_pets) if all_pets else float('nan')
    min_pet = min(all_pets) if all_pets else float('nan')
    
    return {
        "success_rate": success_rate,
        "avg_cav_speed": avg_cav_speed,
        "avg_pet": avg_pet,
        "min_pet": min_pet,
        "total_rounds": total_rounds,
        "successful_rounds": successful_rounds
    }

def main():
    old_highway = "./llm_controller/excel/excel_highway"
    old_intersection = "./llm_controller/excel/excel_intersection"
    old_merge = "./llm_controller/excel/excel_merge"
    
    new_highway = "./results/test/highway-v0/excel"
    new_intersection = "./results/test/intersection-multi-agent-v0/excel"
    new_merge = "./results/test/merge-multi-agent-v0/excel"
    
    scenarios = [
        ("Highway", old_highway, new_highway),
        ("Intersection", old_intersection, new_intersection),
        ("Merge", old_merge, new_merge)
    ]
    
    report_lines = []
    report_lines.append("# CP-V2X Version Comparison Report")
    report_lines.append("\nThis report compares the **Old Version** (v2 first draft with sequential reasoning and legacy safety checks) against the **New Version** (v2 parallelized with cache-RAG and refined heading-projection safety shield).\n")
    report_lines.append("| Scenario | Version | Total Rounds | Success Rate | Avg CAV Speed | Avg PET | Min PET |")
    report_lines.append("| :--- | :--- | :--- | :--- | :--- | :--- | :--- |")
    
    print("\nProcessing and comparing data (optimized openpyxl mode)...")
    
    for name, old_path, new_path in scenarios:
        old_res = analyze_folder(old_path)
        new_res = analyze_folder(new_path)
        
        # Format Old results
        if old_res:
            old_pet = f"{old_res['avg_pet']:.2f}s" if old_res['avg_pet'] == old_res['avg_pet'] else "N/A"
            old_min_pet = f"{old_res['min_pet']:.2f}s" if old_res['min_pet'] == old_res['min_pet'] else "N/A"
            report_lines.append(f"| {name} | **Old** | {old_res['total_rounds']} | {old_res['success_rate']:.1f}% | {old_res['avg_cav_speed']:.2f} m/s | {old_pet} | {old_min_pet} |")
        else:
            report_lines.append(f"| {name} | **Old** | N/A | N/A | N/A | N/A | N/A |")
            
        # Format New results
        if new_res:
            new_pet = f"{new_res['avg_pet']:.2f}s" if new_res['avg_pet'] == new_res['avg_pet'] else "N/A"
            new_min_pet = f"{new_res['min_pet']:.2f}s" if new_res['min_pet'] == new_res['min_pet'] else "N/A"
            report_lines.append(f"| {name} | **New (v2)** | {new_res['total_rounds']} | {new_res['success_rate']:.1f}% | {new_res['avg_cav_speed']:.2f} m/s | {new_pet} | {new_min_pet} |")
        else:
            report_lines.append(f"| {name} | **New (v2)** | N/A | N/A | N/A | N/A | N/A |")
            
    # Write report file
    report_content = "\n".join(report_lines)
    artifact_path = "c:/Users/user/.gemini/antigravity-ide/brain/7a7c0d6b-0ff0-4530-b517-3ad83d5b256f/results_comparison.md"
    os.makedirs(os.path.dirname(artifact_path), exist_ok=True)
    with open(artifact_path, "w", encoding="utf-8") as f:
        f.write(report_content)
        
    print("\n" + "="*80)
    print("                      CP-V2X VERSION COMPARISON REPORT")
    print("="*80)
    print(f"{'Scenario':<15} | {'Version':<10} | {'Rounds':<8} | {'Success%':<10} | {'Avg CAV Speed':<15} | {'Avg PET':<8} | {'Min PET':<8}")
    print("-"*80)
    
    for name, old_path, new_path in scenarios:
        old_res = analyze_folder(old_path)
        new_res = analyze_folder(new_path)
        if old_res:
            old_pet_s = f"{old_res['avg_pet']:.2f}s" if old_res['avg_pet'] == old_res['avg_pet'] else "N/A"
            old_min_s = f"{old_res['min_pet']:.2f}s" if old_res['min_pet'] == old_res['min_pet'] else "N/A"
            print(f"{name:<15} | {'Old':<10} | {old_res['total_rounds']:<8} | {old_res['success_rate']:.1f}% | {old_res['avg_cav_speed']:.2f} m/s     | {old_pet_s:<8} | {old_min_s:<8}")
        if new_res:
            new_pet_s = f"{new_res['avg_pet']:.2f}s" if new_res['avg_pet'] == new_res['avg_pet'] else "N/A"
            new_min_s = f"{new_res['min_pet']:.2f}s" if new_res['min_pet'] == new_res['min_pet'] else "N/A"
            print(f"{name:<15} | {'New (v2)':<10} | {new_res['total_rounds']:<8} | {new_res['success_rate']:.1f}% | {new_res['avg_cav_speed']:.2f} m/s     | {new_pet_s:<8} | {new_min_s:<8}")
    print("="*80)
    print(f"\nReport successfully saved as artifact: {artifact_path}")

if __name__ == "__main__":
    main()
